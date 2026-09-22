# import os
# import time
# import requests
# import pandas as pd
# from pathlib import Path
# from bs4 import BeautifulSoup
# from openpyxl import load_workbook

# # ==========================================
# # CONFIGURATION
# # ==========================================
# LOGIN_URL = "https://myrtpos.com/newbdi/index.fwx"
# REPORT_URL = "https://myrtpos.com/newbdi/Callidus_Recon_Detail.fwx"

# # ==========================================
# # 1. AUTHENTICATION
# # ==========================================
# def login(username, password, log_fn=None):
#     if log_fn:
#         log_fn("Authenticating with portal...")
#     session = requests.Session()
    
#     initial_response = session.get(LOGIN_URL, timeout=30)
#     soup = BeautifulSoup(initial_response.text, 'html.parser')
#     session_id_input = soup.find('input', {'name': 'secSessionID'})
    
#     if not session_id_input:
#         raise ConnectionError("Could not reach the login page.")
        
#     login_payload = {
#         "secSessionID": session_id_input['value'],
#         "secUserID": username, 
#         "secPassword": password,
#         "secSaveCookie": "1"
#     }
    
#     login_response = session.post(LOGIN_URL, data=login_payload, timeout=30)
#     if "secUserID" in login_response.text:
#         raise PermissionError("Login failed! Verify credentials in .env.")
    
#     if log_fn:
#         log_fn("Login successful.")
#     time.sleep(2) 
#     return session

# # ==========================================
# # 2. PROCESS & APPEND DATA
# # ==========================================
# def process_and_append(raw_file_path, store_name, store_id, start_date, end_date, master_file_path, log_fn=None):
#     try:
#         df = pd.read_excel(raw_file_path, engine="xlrd")
        
#         if 'compmrc' in df.columns:
#             df['compmrc'] = df['compmrc'].fillna(0)
        
#         df.insert(0, 'Store ID', store_id)
#         df.insert(1, 'Store Name', store_name)
#         df.insert(2, 'Report Start', start_date)
#         df.insert(3, 'Report End', end_date)
        
#         if not os.path.exists(master_file_path):
#             df.to_excel(master_file_path, index=False, sheet_name="Recon Data")
#             if log_fn:
#                 log_fn(f"[+] Created Master File with {len(df)} rows for {store_name}.")
#         else:
#             wb = load_workbook(master_file_path)
#             ws = wb["Recon Data"] if "Recon Data" in wb.sheetnames else wb.active
            
#             start_row = ws.max_row + 1
#             rows_to_append = df.values.tolist()
            
#             for r_idx, row_data in enumerate(rows_to_append, start=start_row):
#                 for c_idx, value in enumerate(row_data, start=1):
#                     ws.cell(row=r_idx, column=c_idx, value=value)
                    
#             wb.save(master_file_path)
#             if log_fn:
#                 log_fn(f"[+] Appended {len(df)} rows to Master File for {store_name}.")
            
#     except Exception as e:
#         if log_fn:
#             log_fn(f"⚠️ Error processing data for {store_name}: {e}")

# # ==========================================
# # 3. MAIN EXECUTION ENGINE
# # ==========================================
# def execute_recon(username, password, stores, start_date, end_date, workspace_dir: Path, log_fn=None):
#     master_file_path = workspace_dir / f"Callidus_Recon_Master_{start_date.replace('/', '-')}.xlsx"
    
#     if log_fn:
#         log_fn(f"Beginning execution for {len(stores)} stores...")
    
#     session = login(username, password, log_fn)

#     for index, store in enumerate(stores, start=1):
#         store_id = store["id"]
#         safe_store_name = "".join(c for c in store["name"] if c.isalnum() or c in (' ', '-', '_')).strip()
        
#         if log_fn:
#             log_fn(f"[{index}/{len(stores)}] Fetching: {safe_store_name}")
        
#         payload = {
#             "frmMarketID": "", 
#             "frmRegionID": "",
#             "frmStateID": "",
#             "frmStoreType": "",
#             "frmStore": store_id,
#             "frmStart": start_date,
#             "frmEnd": end_date,
#             "frmDateType": "T",
#             "btnExcel": "Export to Excel"
#         }
        
#         max_attempts = 3
#         success = False
#         full_path = workspace_dir / f"temp_{safe_store_name}.xls"
        
#         prev_size = -1
#         prev_rows = -1
        
#         for attempt in range(1, max_attempts + 1):
#             try:
#                 if log_fn:
#                     log_fn(f"➔ Attempt {attempt} of {max_attempts}...")
#                 response = session.post(REPORT_URL, data=payload, stream=True, timeout=60)
                
#                 if "secUserID" in response.text:
#                     if log_fn:
#                         log_fn("⚠️ Session expired. Re-authenticating...")
#                     session = login(username, password, log_fn)
#                     raise ConnectionError("Session dropped.")
                
#                 with open(full_path, "wb") as f:
#                     for chunk in response.iter_content(chunk_size=8192):
#                         f.write(chunk)
                
#                 current_size = full_path.stat().st_size
                
#                 try:
#                     df_check = pd.read_excel(full_path, engine="xlrd")
#                     current_rows = len(df_check)
#                 except Exception:
#                     current_rows = 0

#                 if attempt > 1:
#                     if current_size == prev_size and current_rows == prev_rows and current_rows > 0:
#                         if log_fn:
#                             log_fn(f"✅ Stable download verified! ({current_rows} rows)")
#                         success = True
#                         break
#                     else:
#                         if log_fn:
#                             log_fn(f"Mismatch (Rows: {current_rows} vs {prev_rows}). Retrying...")
                
#                 prev_size = current_size
#                 prev_rows = current_rows
                
#                 if attempt < max_attempts:
#                     time.sleep(3 * attempt)
                    
#             except Exception as e:
#                 if log_fn:
#                     log_fn(f"⚠️ Error on attempt {attempt}: {e}")
#                 time.sleep(3)
        
#         if success:
#             process_and_append(full_path, safe_store_name, store_id, start_date, end_date, master_file_path, log_fn)
#             if os.path.exists(full_path):
#                 os.remove(full_path)
#         else:
#             if log_fn:
#                 log_fn(f"❌ Failed to achieve stable download after {max_attempts} attempts.")
            
#         time.sleep(1)
        
#     return master_file_path
import os
import time
import requests
import pandas as pd
from pathlib import Path
from bs4 import BeautifulSoup
from openpyxl import load_workbook

# ==========================================
# CONFIGURATION
# ==========================================
LOGIN_URL = "https://myrtpos.com/newbdi/index.fwx"
REPORT_URL = "https://myrtpos.com/newbdi/Callidus_Recon_Detail.fwx"

# ==========================================
# 1. AUTHENTICATION
# ==========================================
def login(username, password, log_fn=None):
    if log_fn:
        log_fn("Authenticating with portal...")
    
    session = requests.Session()
    initial_response = session.get(LOGIN_URL, timeout=30)
    soup = BeautifulSoup(initial_response.text, 'html.parser')
    session_id_input = soup.find('input', {'name': 'secSessionID'})
    
    if not session_id_input:
        raise ConnectionError("Could not reach the login page.")
        
    login_payload = {
        "secSessionID": session_id_input['value'],
        "secUserID": username, 
        "secPassword": password,
        "secSaveCookie": "1"
    }
    
    login_response = session.post(LOGIN_URL, data=login_payload, timeout=30)
    if "secUserID" in login_response.text:
        raise PermissionError("Login failed! Verify credentials in .env.")
    
    if log_fn:
        log_fn("Login successful.")
    time.sleep(2) 
    return session

# ==========================================
# 2. PROCESS & APPEND DATA
# ==========================================
def process_and_append(raw_file_path, store_name, store_id, start_date, end_date, master_file_path, log_fn=None):
    try:
        df = pd.read_excel(raw_file_path, engine="xlrd")
        
        if 'compmrc' in df.columns:
            df['compmrc'] = df['compmrc'].fillna(0)
        
        df.insert(0, 'Store ID', store_id)
        df.insert(1, 'Store Name', store_name)
        df.insert(2, 'Report Start', start_date)
        df.insert(3, 'Report End', end_date)
        
        if not os.path.exists(master_file_path):
            df.to_excel(master_file_path, index=False, sheet_name="Recon Data")
            if log_fn:
                log_fn(f"[+] Created Master File for {store_name}.")
        else:
            wb = load_workbook(master_file_path)
            ws = wb["Recon Data"] if "Recon Data" in wb.sheetnames else wb.active
            
            start_row = ws.max_row + 1
            rows_to_append = df.values.tolist()
            
            for r_idx, row_data in enumerate(rows_to_append, start=start_row):
                for c_idx, value in enumerate(row_data, start=1):
                    ws.cell(row=r_idx, column=c_idx, value=value)
                    
            wb.save(master_file_path)
            if log_fn:
                log_fn(f"[+] Appended data for {store_name}.")
            
    except Exception as e:
        if log_fn:
            log_fn(f"⚠️ Error processing data for {store_name}: {e}")

# ==========================================
# 3. MAIN EXECUTION ENGINE
# ==========================================
# Note the addition of `progress_fn=None` in the arguments below
def execute_recon(username, password, stores, start_date, end_date, workspace_dir: Path, log_fn=None, progress_fn=None):
    master_file_path = workspace_dir / f"Callidus_Recon_Master_{start_date.replace('/', '-')}.xlsx"
    total_stores = len(stores)
    
    if progress_fn:
        progress_fn(0, total_stores, "Authenticating...")
    
    session = login(username, password, log_fn)

    for index, store in enumerate(stores, start=1):
        store_id = store["id"]
        safe_store_name = "".join(c for c in store["name"] if c.isalnum() or c in (' ', '-', '_')).strip()
        
        # Trigger the frontend progress bar update
        if progress_fn:
            progress_fn(index, total_stores, f"Fetching {safe_store_name}...")
        
        payload = {
            "frmMarketID": "", 
            "frmRegionID": "",
            "frmStateID": "",
            "frmStoreType": "",
            "frmStore": store_id,
            "frmStart": start_date,
            "frmEnd": end_date,
            "frmDateType": "T",
            "btnExcel": "Export to Excel"
        }
        
        max_attempts = 3
        success = False
        full_path = workspace_dir / f"temp_{safe_store_name}.xls"
        
        prev_size = -1
        prev_rows = -1
        
        for attempt in range(1, max_attempts + 1):
            try:
                if log_fn: log_fn(f"➔ Attempt {attempt} of {max_attempts}...")
                
                response = session.post(REPORT_URL, data=payload, stream=True, timeout=60)
                
                if "secUserID" in response.text:
                    if log_fn: log_fn("⚠️ Session expired. Re-authenticating...")
                    session = login(username, password, log_fn)
                    raise ConnectionError("Session dropped.")
                
                with open(full_path, "wb") as f:
                    for chunk in response.iter_content(chunk_size=8192):
                        f.write(chunk)
                
                current_size = full_path.stat().st_size
                
                try:
                    df_check = pd.read_excel(full_path, engine="xlrd")
                    current_rows = len(df_check)
                except Exception:
                    current_rows = 0

                if attempt > 1:
                    if current_size == prev_size and current_rows == prev_rows and current_rows > 0:
                        if log_fn: log_fn(f"✅ Stable download verified! ({current_rows} rows)")
                        success = True
                        break
                    else:
                        if log_fn: log_fn(f"Mismatch (Rows: {current_rows} vs {prev_rows}). Retrying...")
                
                prev_size = current_size
                prev_rows = current_rows
                
                if attempt < max_attempts:
                    time.sleep(3 * attempt)
                    
            except Exception as e:
                if log_fn: log_fn(f"⚠️ Error on attempt {attempt}: {e}")
                time.sleep(3)
        
        if success:
            process_and_append(full_path, safe_store_name, store_id, start_date, end_date, master_file_path, log_fn)
            if os.path.exists(full_path):
                os.remove(full_path)
        else:
            if log_fn: log_fn(f"❌ Failed to achieve stable download after {max_attempts} attempts.")
            
        time.sleep(1)
        
    return master_file_path